from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

wb = Workbook()
INK, ACC = '1B1F27', 'C2410C'
hf = Font(name='Arial', bold=True, color='FFFFFF', size=10)
hfill = PatternFill('solid', fgColor=INK)
base = Font(name='Arial', size=10)
bold = Font(name='Arial', size=10, bold=True)
sim = PatternFill('solid', fgColor='FDE7DA')
box = Border(bottom=Side(style='thin', color='D9DCE3'))
BANNER = ('SIMULATED DATA. No participant has used this MVP yet. Every user, quote and number '
          'below is illustrative and must be replaced with real observations.')


def sheet(ws, title, headers, rows, widths, banner=True):
    ws.title = title
    r = 1
    if banner:
        c = ws.cell(1, 1, BANNER)
        c.font = Font(name='Arial', bold=True, color=ACC, size=10)
        c.fill = sim
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        r = 3
    for j, h in enumerate(headers, 1):
        c = ws.cell(r, j, h)
        c.font, c.fill = hf, hfill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(rows, r + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font, c.border = base, box
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = ws.cell(r + 1, 1)
    return r


# ---------- Read Me
ws = wb.active
ws.title = 'Read Me'
lines = [
    ('SLIDE 8: V1 USABILITY TEST REPORT (SIMULATED)', Font(name='Arial', bold=True, size=14, color=INK)),
    (BANNER, Font(name='Arial', bold=True, color=ACC)),
    ('', base),
    ('Purpose', bold),
    ('Shows what a round of 3 usability sessions on the V1 adaptive memory-retrieval MVP would record, '
     'what it would teach, and what changes in V2. The layout is ready for real sessions; the data is worked example data.', base),
    ('', base),
    ('How to convert to real results', bold),
    ('1. Run 3 sessions with Segment T users (they expect a specific photo to exist, remember it imprecisely, and keep trying).  '
     '2. Overwrite the shaded cells in Participants and Sessions.  3. Metrics recalculates.  '
     '4. Rewrite Learnings and V1 to V2 from what you saw, then remove every "Simulated" label.', base),
    ('', base),
    ('Sheets', bold),
    ('Participants | Tasks | Sessions | Metrics | Learnings | V1 to V2', base),
    ('', base),
    ('What is NOT simulated', bold),
    ('The Tasks sheet uses the six real memory scenarios and 72 photos in the MVP. Learning L1 (filter options are never shown) '
     'comes from reading the V1 source code, not from a session.', base),
    ('', base),
    ('Limits', bold),
    ('3 users and 6 sessions, so each session moves a percentage by about 17 points. V1 does not log time, inspections or '
     'recovery steps, so a facilitator would record them by hand.', base),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 120
ws.sheet_view.showGridLines = False

# ---------- Participants
P = [('U1', 'SIMULATED', 'Product designer, 29', 'Google Photos, daily',
      'Hunts for a photo in 2+ attempts each month; switches between search and scrolling', 'Segment T: effortful-path retriever'),
     ('U2', 'SIMULATED', 'Marketing manager, 34', 'Google Photos, weekly',
      'Remembers place and people but not dates; scrolls the timeline when search fails', 'Segment T: effortful-path retriever'),
     ('U3', 'SIMULATED', 'Graduate student, 24', 'Google Photos, weekly',
      'Uses vague search words, retries with new words, gives up after several misses', 'Segment T: effortful-path retriever')]
r = sheet(wb.create_sheet(), 'Participants',
          ['User', 'Status', 'Profile', 'Photos use', 'Retrieval behaviour (screener)', 'Segment match'],
          P, [8, 14, 26, 22, 64, 34])
for i in range(r + 1, r + 4):
    for j in (3, 4, 5):
        wb['Participants'].cell(i, j).fill = sim

# ---------- Tasks
T = [('T1', 'Beach cafe with my cousin around sunset in Goa', 'Goa, cousin, beach cafe, sunset', 'Exact date, file name, keyword', 'Real scenario in the MVP (12 photos)'),
     ('T2', 'Hiking in the monsoon near Maharashtra mountains', 'Maharashtra, rain, hiking', 'Exact date, trail name, keyword', 'Real scenario in the MVP (12 photos)'),
     ('T3', 'Late night beach party with friends in Goa', 'Goa, friends, night, party, guitar', 'Exact date, album, keyword', 'Real scenario in the MVP (12 photos)'),
     ('T4', 'That hidden waterfall we found in the jungle', 'Waterfall, jungle, hidden', 'Place name, date, keyword', 'Real scenario in the MVP (12 photos)'),
     ('T5', 'The houseboat trip we took in Kerala with family', 'Kerala, houseboat, family', 'Exact date, album, keyword', 'Real scenario in the MVP (12 photos)'),
     ('T6', 'Watching the sunset from a hammock', 'Sunset, hammock', 'Place, date, keyword', 'Real scenario in the MVP (12 photos)')]
sheet(wb.create_sheet(), 'Tasks',
      ['Task', 'Memory the user is given (scenario card)', 'Clues they may use', 'Withheld from the user', 'Source'],
      T, [8, 52, 40, 32, 36], banner=False)

# ---------- Sessions
H = ['Session', 'User', 'Task', 'Query typed (simulated)', 'Clues read correctly', 'First attempt found it',
     'Recovery used (count)', 'Strategy switches', 'Candidates inspected', 'Time to retrieval (s)',
     'Found intended photo', 'Confidence (1-5)', 'Observation (simulated)']
S = [('S1', 'U1', 'T1', 'beach cafe with my cousin sunset goa', 'Yes', 'Yes', 0, 0, 4, 48, 'Yes', 4,
      'Read the clue chips and picked the second photo quickly.'),
     ('S2', 'U1', 'T4', 'hidden waterfall trek with friends', 'Partly', 'No', 3, 2, 7, 142, 'Yes', 3,
      '"friends" became a hard filter and removed every waterfall photo. Recovery showed 10 categories; user said "I do not know which one matters".'),
     ('S3', 'U2', 'T3', 'night party goa beach with my college friends', 'Yes', 'No', 1, 1, 9, 96, 'Yes', 4,
      'First set looked right but not the photo with the guitar; found it through Entertainment > Guitar.'),
     ('S4', 'U2', 'T5', 'boat ride kerala backwaters with my parents', 'Partly', 'Yes', 0, 0, 3, 41, 'Yes', 5,
      '"parents" was not read as a clue but nothing depended on it.'),
     ('S5', 'U3', 'T2', 'trek near lonavala in rain', 'No', 'No', 2, 3, 8, 210, 'No', 2,
      'Lonavala not recognised as Maharashtra, so no location clue and weak matches. Gave up after two recovery attempts.'),
     ('S6', 'U3', 'T6', 'hammock sunset', 'Yes', 'No', 1, 1, 11, 118, 'Yes', 3,
      'Many similar sunsets from different places. Asked "can I narrow by place?" but no filters were on screen.')]
r = sheet(wb.create_sheet(), 'Sessions', H, S, [9, 7, 7, 40, 14, 14, 14, 12, 13, 14, 14, 12, 72])
a, b = r + 1, r + len(S)
for i in range(a, b + 1):
    for j in range(4, 14):
        wb['Sessions'].cell(i, j).fill = sim


def rng(col):
    return f"Sessions!${col}${a}:${col}${b}"


# ---------- Metrics
ws = wb.create_sheet('Metrics')
ws['A1'] = BANNER
ws['A1'].font = Font(name='Arial', bold=True, color=ACC)
ws['A1'].fill = sim
ws.merge_cells('A1:F1')
for j, h in enumerate(['Metric', 'Overall', 'U1', 'U2', 'U3', 'How it is calculated'], 1):
    c = ws.cell(3, j, h)
    c.font, c.fill = hf, hfill
U = ('U1', 'U2', 'U3')
rows = [
    ('Retrieval success (% of sessions)', f'=COUNTIF({rng("K")},"Yes")/COUNTA({rng("K")})',
     [f'=COUNTIFS({rng("B")},"{u}",{rng("K")},"Yes")/COUNTIF({rng("B")},"{u}")' for u in U],
     'Sessions where the user identified the intended photo'),
    ('First-attempt success (% of sessions)', f'=COUNTIF({rng("F")},"Yes")/COUNTA({rng("F")})',
     [f'=COUNTIFS({rng("B")},"{u}",{rng("F")},"Yes")/COUNTIF({rng("B")},"{u}")' for u in U],
     'Found without pressing "none of these"'),
    ('Recovery rate (% of failed first attempts that ended in success)',
     f'=COUNTIFS({rng("F")},"No",{rng("K")},"Yes")/COUNTIF({rng("F")},"No")',
     [f'=IFERROR(COUNTIFS({rng("B")},"{u}",{rng("F")},"No",{rng("K")},"Yes")/COUNTIFS({rng("B")},"{u}",{rng("F")},"No"),"n/a")' for u in U],
     'Failed first attempts that still found the photo'),
    ('Median time to retrieval, successful sessions (s)', None, ['', '', ''],
     'Array formula over successful sessions'),
    ('Average strategy switches per session', f'=AVERAGE({rng("H")})',
     [f'=AVERAGEIF({rng("B")},"{u}",{rng("H")})' for u in U], 'Changes of approach during a session'),
    ('Average candidates inspected per session', f'=AVERAGE({rng("I")})',
     [f'=AVERAGEIF({rng("B")},"{u}",{rng("I")})' for u in U], 'Proxy for manual browsing'),
    ('Average recognition confidence (1-5)', f'=AVERAGE({rng("L")})',
     [f'=AVERAGEIF({rng("B")},"{u}",{rng("L")})' for u in U], 'Self-reported after each session'),
    ('Clues read correctly (% of sessions)', f'=COUNTIF({rng("E")},"Yes")/COUNTA({rng("E")})',
     [f'=COUNTIFS({rng("B")},"{u}",{rng("E")},"Yes")/COUNTIF({rng("B")},"{u}")' for u in U],
     'Express quality: every key clue extracted'),
]
for i, (m, o, ps, how) in enumerate(rows, 4):
    ws.cell(i, 1, m).font = bold
    ws.cell(i, 6, how).font = base
    if o:
        ws.cell(i, 2, o)
    for j, p in enumerate(ps, 3):
        ws.cell(i, j, p)
    for j in range(2, 6):
        ws.cell(i, j).font = base
        ws.cell(i, j).number_format = '0%' if '%' in m else '0.0'
ws['B7'] = ArrayFormula('B7', f'=MEDIAN(IF({rng("K")}="Yes",{rng("J")}))')
ws['B7'].number_format = '0'
ws['A13'] = 'With 3 users each session is worth about 17 percentage points. Read these as signals to check, not rates.'
ws['A13'].font = Font(name='Arial', italic=True, color='8A90A0', size=9)
for j, w in enumerate([62, 12, 10, 10, 10, 56], 1):
    ws.column_dimensions[L(j)].width = w

# ---------- Learnings
Lr = [('L1', 'Recognition layer', 'Code check, not simulated: the filter options are computed after each search but never rendered, so users cannot narrow by place or year.', 'S6 (asked for a place filter)', 'High', 'Render the filters, and recommend the one most likely to help instead of listing all.', 'Code check'),
      ('L2', 'Express', 'Plain-word rules miss real phrasing ("parents", "Lonavala", "trek"). A wrong clue can also hard-filter everything out ("friends" in S2).', 'S2, S4, S5', 'High', 'Use an LLM to read the memory; treat uncertain clues as soft signals and confirm before filtering.', 'Simulated'),
      ('L3', 'Recover', 'The recovery list has 6 to 10 unranked categories. Users did not know which one mattered.', 'S2, S3, S5', 'High', 'Offer 2 or 3 ranked recovery actions based on what was already tried, and say why.', 'Simulated'),
      ('L4', 'Match', 'When the clues are read right (S1, S4), users finish in under a minute. The loop works.', 'S1, S4', 'Info', 'Keep the loop; fix Express and Recover first.', 'Simulated'),
      ('L5', 'Match', 'With several similar photos, users inspected 9 to 11 before finding the intended one.', 'S3, S6', 'Medium', 'Improve ranking so the intended photo is in the first 3.', 'Simulated'),
      ('L6', 'Measurement', 'V1 records no time, inspections or recovery steps; a facilitator logged them by hand.', 'All', 'Medium', 'Log session events in the app.', 'Code check')]
sheet(wb.create_sheet(), 'Learnings',
      ['ID', 'Journey stage', 'Learning', 'Evidence', 'Severity', 'What to change in V2', 'Basis'],
      Lr, [6, 18, 72, 26, 10, 62, 14])

# ---------- V1 to V2
V = [('Natural-language memory input', 'Can users describe the photo naturally?', 'Users described photos naturally, but 3 of 6 sessions had a clue missed or misread.', 'Which clues are easiest and hardest to interpret?', 'Place names and relationships outside the rule list fail.', 'Improve context extraction (LLM, soft filters)', 'Simulated'),
     ('Contextual candidate retrieval', 'Are returned candidates relevant enough?', 'Good when clues are right (S1, S4); poor when one clue is wrong (S2, S5).', 'Which clue combinations work?', 'One wrong hard filter can empty the set.', 'Improve candidate generation; match softly', 'Simulated'),
     ('Dynamic candidate filters', 'Can users recognise the photo faster?', 'Not testable in V1: filters are not shown.', 'Which filters help recognition?', 'A user asked for place narrowing (S6).', 'Show and personalise filters', 'Code check + simulated'),
     ('Guided recovery', 'What do users do after the first attempt fails?', '4 of 4 failed first attempts used recovery; 3 of 4 then found the photo.', 'Which suggestions help users progress?', 'A long unranked list did not help; specific signals (Guitar, Waterfall) did.', 'Rank recovery actions by context', 'Simulated')]
sheet(wb.create_sheet(), 'V1 to V2',
      ['V1 feature', 'Test / observation', 'What we saw', 'Learning to validate', 'What it means', 'V2 direction', 'Basis'],
      V, [28, 32, 48, 32, 44, 38, 20])

wb.save('Slide08_V1_Usability_Test_Report_SIMULATED.xlsx')
