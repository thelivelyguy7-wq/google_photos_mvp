import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

SRC = r'C:\Users\ADMIN\Downloads\completed mvp testing.xlsx'
OUT = 'MVP_testing_SIMULATED_corrected.xlsx'
wb = openpyxl.load_workbook(SRC)

INK, ACC = '1B1F27', 'C2410C'
hf = Font(name='Arial', bold=True, color='FFFFFF', size=10)
hfill = PatternFill('solid', fgColor=INK)
base = Font(name='Arial', size=10)
bold = Font(name='Arial', size=10, bold=True)
sim = PatternFill('solid', fgColor='FDE7DA')
box = Border(bottom=Side(style='thin', color='D9DCE3'))
BANNER = ('SIMULATED EXAMPLE. Illustrative test format, not real user data. No V2 has been built or tested. '
          'Rows marked "Code check" come from reading the V1 source code.')


def cell(ws, ref, v, font=base, fill=None, wrap=True):
    c = ws[ref]
    c.value = v
    c.font = font
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    c.border = box
    if fill: c.fill = fill
    return c


def header(ws, row, col, text):
    c = ws.cell(row, col, text)
    c.font, c.fill = hf, hfill
    c.alignment = Alignment(wrap_text=True, vertical='center')


for name in ['Participants', 'Sessions', 'Metrics', 'Learnings', 'V1 to V2', 'V2 Changes', 'V2 Sessions',
             'V2 Participants', 'V1 vs V2 Metrics', 'V2 Learnings and Next']:
    ws = wb[name]
    ws['A1'].value = BANNER
    ws['A1'].font = Font(name='Arial', bold=True, color=ACC, size=10)
    ws['A1'].fill = sim

# ---------------- Read Me
rm = wb['Read Me']
rm['A1'] = 'SLIDE 8: V1 TO V2 USABILITY TEST REPORT (SIMULATED EXAMPLE)'
rm['A2'] = BANNER
rm['A2'].font = Font(name='Arial', bold=True, color=ACC)
rm['A2'].alignment = Alignment(wrap_text=True, vertical='top')
rm['A5'] = ('Shows the format of a round of 3 usability sessions on the V1 adaptive memory-retrieval MVP, what it would teach, '
            'and a simulated V2 round. The layout is ready for real sessions; the data is worked example data.')
rm['A11'] = ('Real vs Simulated | Participants | Tasks | Sessions | Metrics | Learnings | V1 to V2 | V2 Changes | '
             'V2 Participants | V2 Sessions | V1 vs V2 Metrics | V2 Learnings and Next')
rm['A13'] = 'What is NOT simulated'
rm['A13'].font = bold
rm['A14'] = ('The six task scenarios and the 72 photos are the MVP\'s own data. Findings marked "Code check" (filters never shown, '
             'only 6 of 12 photos shown, unranked recovery, no logging) come from reading the V1 source code. '
             'See the Real vs Simulated sheet.')
rm['A17'] = ('3 users and 6 sessions per round, so one session moves a percentage by about 17 points. '
             'V1 does not log time, inspections or recovery steps, so a facilitator would record them by hand.')
rm['A19'] = 'V2 ROUND (SIMULATED)'
rm['A19'].font = Font(name='Arial', bold=True, color=ACC)
rm['A20'] = ('Added sheets: V2 Changes | V2 Participants | V2 Sessions | V1 vs V2 Metrics | V2 Learnings and Next. '
             'The V1 to V2 sheet has V2 built, V2 result, Status and Basis columns. All V2 content is simulated: no V2 was built or tested.')
for r in (5, 11, 14, 17, 20):
    rm.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')

# ---------------- Real vs Simulated (new)
rs = wb.create_sheet('Real vs Simulated', 1)
rs['A1'] = 'What in this workbook is real, and what is illustrative'
rs['A1'].font = Font(name='Arial', bold=True, size=12, color=INK)
for j, h in enumerate(['Item', 'Status', 'Source or note'], 1):
    header(rs, 3, j, h)
RS = [
    ('Six memory scenarios and 72 photos', 'Real', 'MVP data (src/mockData.js, public/google)'),
    ('Filter options computed after each search but never displayed', 'Real (code check)', 'src/App.jsx: availableFilters is set, not rendered'),
    ('Only the first 6 of 12 photos in the best-matching group are shown', 'Real (code check)', 'src/App.jsx: candidates.slice(0, 6)'),
    ('Recovery is an unranked list of 6 to 10 categories', 'Real (code check)', 'src/App.jsx: recovery taxonomies'),
    ('Clues are read by fixed keyword rules; every extracted clue is a hard filter', 'Real (code check)', 'src/App.jsx: mockContextExtractor; src/mockData.js: filterByMetadata'),
    ('No in-app logging of time, inspections or recovery steps', 'Real (code check)', 'The success screen shows a label only'),
    ('Task links to discovery patterns', 'From discovery dataset', 'Slide 3 counts. The discovery dataset is itself simulated, per the deck'),
    ('Participants U1-U6, sessions S1-S12, quotes, times, confidence', 'Simulated', 'Invented to show the test format'),
    ('All metric values', 'Simulated inputs, real formulas', 'Formulas recalculate from the Sessions sheets'),
    ('V2 features (LLM clue reading, filter chips, ranked recovery, show all, logging)', 'Simulated', 'Not built. They are proposals'),
    ('V2 results and the V1 vs V2 comparison', 'Simulated', 'No V2 was tested'),
    ('Priorities and owners in V2 Learnings and Next', 'Proposal', 'A suggested order, not a decision'),
]
for i, row in enumerate(RS, 4):
    for j, v in enumerate(row, 1):
        c = rs.cell(i, j, v)
        c.font = bold if j == 2 else base
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = box
        if row[1].startswith('Simulated') or row[1] == 'Proposal': c.fill = sim
for j, w in enumerate([70, 30, 70], 1):
    rs.column_dimensions[L(j)].width = w
rs.sheet_view.showGridLines = False

# ---------------- Participants / V2 Participants status
for r in range(4, 7):
    wb['Participants'].cell(r, 2).value = 'Simulated'
    wb['V2 Participants'].cell(r, 2).value = 'Simulated'

# ---------------- Tasks: discovery pattern and intended photo
tk = wb['Tasks']
header(tk, 1, 6, 'Discovery pattern (Slide 3)')
header(tk, 1, 7, 'Intended photo ID (chosen for the test)')
TK = {2: ('Travel / place photo; holds people and approximate time, lacks exact date (Slide 3)', 101),
      3: ('Travel / place photo; holds activity, lacks a name or keyword (Slide 3)', 118),
      4: ('Event / social photo; holds people and approximate time, lacks exact date (Slide 3)', 130),
      5: ('Travel / place photo; holds visual appearance, lacks a name or keyword (Slide 3)', 138),
      6: ('Travel / place photo with family; holds people, lacks exact date (Slide 3)', 150),
      7: ('Visual-appearance memory; holds visual appearance, lacks a keyword (Slide 3)', 162)}
for r, (pat, pid) in TK.items():
    cell(tk, f'F{r}', pat)
    cell(tk, f'G{r}', pid)
for col, w in zip('FG', (62, 22)):
    tk.column_dimensions[col].width = w
tk.column_dimensions['E'].width = 30

# ---------------- V1 Sessions
ss = wb['Sessions']
ss['M7'].value = '"parents" was not read as a clue but nothing depended on it.'
ss['E8'].value = 'Partly'
ss['M8'].value = ('"trek" and "rain" matched, so hiking photos appeared, but only 6 of 12 in the group were shown and the intended '
                  'photo was in the hidden half. Gave up after two recovery attempts. Lonavala was not read as Maharashtra, '
                  'which did not change the results.')
ss['M6'].value = ('First 6 of 12 photos looked right but not the one with the guitar; found it through Entertainment > Guitar.')
ss['M9'].value = ('Several near-identical Goa sunsets from two scenarios. Asked "can I narrow this down?" but no filters were on screen.')
for j, h in zip((14, 15, 16), ('Intended photo ID', 'Photo picked ID', 'Participant quote (simulated)')):
    header(ss, 3, j, h)
picked1 = {4: 101, 5: 138, 6: 130, 7: 150, 8: '-', 9: 162}
quote1 = {4: '"That\'s the one, second photo."', 5: '"I don\'t know which of these matters."',
          6: '"Right night, but not the photo with the guitar."', 7: '"Easy, that\'s the houseboat."',
          8: '"I can\'t tell these hiking photos apart."', 9: '"Can I narrow this down somehow?"'}
for r in range(4, 10):
    cell(ss, f'N{r}', f'=VLOOKUP(C{r},Tasks!$A$2:$G$7,7,FALSE)', fill=sim)
    cell(ss, f'O{r}', picked1[r], fill=sim)
    cell(ss, f'K{r}', f'=IF(O{r}=N{r},"Yes","No")', fill=sim)
    cell(ss, f'P{r}', quote1[r], fill=sim)
    ss[f'E{r}'].fill = ss[f'M{r}'].fill = sim
for col, w in zip('NOP', (14, 14, 44)):
    ss.column_dimensions[col].width = w
ss.column_dimensions['M'].width = 72

# ---------------- V2 Sessions
v2 = wb['V2 Sessions']
for j, h in zip((15, 16, 17), ('Intended photo ID', 'Photo picked ID', 'Participant quote (simulated)')):
    header(v2, 3, j, h)
picked2 = {4: 101, 5: 138, 6: 130, 7: 150, 8: 118, 9: 104}
quote2 = {4: '"Oh, that was quick."', 5: '"It kept the waterfalls even though I said friends."',
          6: '"Show all was what I was missing."', 7: '"Straight to it."',
          8: '"Adding morning narrowed it down."', 9: '"That\'s the one."'}
for r in range(4, 10):
    cell(v2, f'O{r}', f'=VLOOKUP(C{r},Tasks!$A$2:$G$7,7,FALSE)', fill=sim)
    cell(v2, f'P{r}', picked2[r], fill=sim)
    cell(v2, f'K{r}', f'=IF(P{r}=O{r},"Yes","No")', fill=sim)
    cell(v2, f'Q{r}', quote2[r], fill=sim)
for col, w in zip('OPQ', (14, 14, 44)):
    v2.column_dimensions[col].width = w

# ---------------- Learnings
ln = wb['Learnings']
ln['C4'].value = ('Code check: the filter options are computed after each search but never rendered, so users cannot narrow the set '
                  'by place or year.')
ln['D4'].value = 'S6 (asked to narrow down)'
ln['C5'].value = ('Plain-word rules miss real phrasing ("parents", "Lonavala"). A wrong clue can also hard-filter everything out '
                  '("friends" in S2).')
ln['C8'].value = ('Only 6 of 12 photos in the best group are shown, and several photos look alike, so users inspected 8 to 11 '
                  'photos or gave up.')
ln['D8'].value = 'S3, S5, S6'
for r, b in zip(range(4, 10), ['Code check', 'Simulated', 'Simulated', 'Simulated', 'Code check + simulated', 'Code check']):
    cell(ln, f'G{r}', b, fill=sim if 'imulated' in b else None)

# ---------------- V1 to V2
vt = wb['V1 to V2']
header(vt, 3, 10, 'Basis')
vt['E6'].value = 'A user asked to narrow down (S6).'
vt['E5'].value = 'One wrong hard filter can empty the set, and only 6 of 12 photos in the best group are shown.'
for r, b in zip(range(4, 10), ['Simulated', 'Code check + simulated', 'Code check + simulated', 'Simulated', 'Simulated', 'Simulated']):
    cell(vt, f'J{r}', b, fill=sim)
    for col in 'GHI':
        vt[f'{col}{r}'].fill = sim
vt.column_dimensions['J'].width = 24

# ---------------- V2 Changes / V2 Learnings basis
ch = wb['V2 Changes']
header(ch, 3, 6, 'Basis')
for r in range(4, 10):
    st = ch[f'E{r}'].value
    ch[f'E{r}'].value = 'Built (simulated)' if st == 'Built' else 'Kept'
    cell(ch, f'F{r}', 'Simulated', fill=sim)
ch.column_dimensions['E'].width = 18
vl = wb['V2 Learnings and Next']
header(vl, 3, 9, 'Basis')
for r in range(4, 9):
    cell(vl, f'I{r}', 'Simulated' if r != 8 else 'Simulated', fill=sim)
vl.column_dimensions['I'].width = 12

# ---------------- V1 vs V2 metrics note
vm = wb['V1 vs V2 Metrics']
vm['A14'].value = ('V1 and V2 used different simulated users on the same 6 tasks (S1-S6 vs S7-S12). With 3 users each session is worth '
                   'about 17 points, so read the pattern, not the exact size.')

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT, wb.sheetnames)
